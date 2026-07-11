"""Generates the blank Excel material-tracking template (Lamela 1-6)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = "Lamela_Material_Tracking.xlsx"

COLUMNS = [
    ("Датум", 14),
    ("Фирма", 28),
    ("Спрат", 12),
    ("Материјал", 26),
    ("Количина", 14),
    ("Забелешка", 30),
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()
wb.remove(wb.active)

for i in range(1, 7):
    ws = wb.create_sheet(title=f"Lamela {i}")

    for col_idx, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    for row in range(2, 202):
        for col_idx in range(1, len(COLUMNS) + 1):
            ws.cell(row=row, column=col_idx).border = BORDER
        if COLUMNS[0][0] == "Датум":
            ws.cell(row=row, column=1).number_format = "dd.mm.yyyy"

wb.save(OUTPUT_PATH)
print(f"Saved {OUTPUT_PATH}")
